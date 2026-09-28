"""Leakage: defect leakage per release, from Jira, as Delivery counts it.

No network.  What is locked here is what decides the figures a manager reads
and the proposals a Key QA reviews: which fixVersions are releases, where
each window starts and ends, which incidents count, the ratios, how an
incident is matched to the TestRail case that should have caught it, what
the AI's answer is allowed to say, and that a review never erases the AI's
proposal.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timezone

import pandas as pd
import pytest

from src import jira_client as jc
from src import leakage as lk
from src import leakage_ai as lai
from src import leakage_export as lx
from src import leakage_match as lm
from src import leakage_store as ls

TODAY = date(2026, 9, 28)


def _v(name, released, when):
    return {"name": name, "released": released, "releaseDate": when}


EE_VERSIONS = [
    _v("EE_SAP_Release_2026Q2.Apr", True, "2026-04-27"),
    _v("EE_SAP_Hotfix1_2026Q2.Apr", True, "2026-05-06"),
    _v("EE_SAP_Release_2026Q2.Jun", True, "2026-06-16"),
    _v("EE MAPP Release 4.38", True, "2026-06-01"),
    _v("EE_SAP_Release_2026Q3.Sept", True, "2026-09-21"),
    _v("EE_SAP_Release_2026Q4.Oct", False, "2026-10-07"),
    _v("EE_SAP_Release_2026Q4.Dec", False, None),
]
EE = next(g for g in lk.GROUPS if g.project == "EE20")


def _incident(key, created, *, rc="Code issue", priority="Medium", status="Done",
              components=(), labels=()):
    return {"key": key, "url": f"https://j/browse/{key}", "type": "Production Incident",
            "summary": "Checkout fails with voucher", "created": created,
            "priority": priority, "status": status, "components": list(components),
            "labels": list(labels), "root_cause": rc, "environment": "PROD",
            "description": "", "steps": "", "actual": "", "expected": "",
            "case_refs": "", "links": [], "resolution": ""}


# ── releases and windows ─────────────────────────────────────────────────────
class TestReleases:
    def test_only_the_main_stream_is_a_release(self):
        """Hotfixes and MAPP versions are other streams: a hotfix between two
        releases must not cut the first one's window short."""
        rels = lk.releases_from_versions(EE_VERSIONS, EE.release_pattern)
        assert [r.name for r in rels] == [
            "EE_SAP_Release_2026Q2.Apr", "EE_SAP_Release_2026Q2.Jun",
            "EE_SAP_Release_2026Q3.Sept", "EE_SAP_Release_2026Q4.Oct",
            "EE_SAP_Release_2026Q4.Dec"]

    def test_every_group_names_real_dashboard_bus(self):
        from src.bu_rules import ALL_RULES
        bus = {r.bu for r in ALL_RULES}
        for g in lk.GROUPS:
            assert set(g.bus) <= bus, g.label

    def test_the_stream_patterns_match_the_names_jira_uses(self):
        """Names read from Jira on 2026-09-28, one per project."""
        names = {"IPXL20": "IPXL_SAP_Release_2026Q3.2", "KRUIDVAT": "KV_SAP_Release_2026Q3.2",
                 "MRN20": "MRN_SAP_Release_2026Q3.1", "TPS20": "TPS_SAP_Release_2026Q3.2",
                 "SD20": "SD_Release_2026Q3.Sep SAP", "EE20": "EE_SAP_Release_2026Q3.Sept"}
        import re
        for g in lk.GROUPS:
            assert re.match(g.release_pattern, names[g.project]), g.project
        kv = next(g for g in lk.GROUPS if g.project == "KRUIDVAT")
        for other in ("TP_SAP_Release_2026.1", "KV_SAP_Release_2026Q4.1 remove",
                      "KV_SAP_HotfixSPA_2026Q2.10"):
            assert not re.match(kv.release_pattern, other), other

    def test_shared_projects_are_declared_groups_not_a_split(self):
        ee = lk.group_for_bu("Drogas")
        assert ee.shared and ee is lk.group_for_bu("Watsons Turkey") is lk.group_for_bu("Watsons Ukraine")
        assert lk.group_for_bu("Trekpleister") is None

    def test_latest_released_is_the_default(self):
        rels = lk.releases_from_versions(EE_VERSIONS, EE.release_pattern)
        assert lk.latest_released(rels).name == "EE_SAP_Release_2026Q3.Sept"


class TestWindows:
    rels = lk.releases_from_versions(EE_VERSIONS, EE.release_pattern)

    def test_a_closed_window_ends_when_the_next_release_shipped(self):
        w = lk.window_for(self.rels, "EE_SAP_Release_2026Q2.Apr", TODAY)
        assert (w.start, w.end, w.open_ended) == (date(2026, 4, 27), date(2026, 6, 16), False)
        assert w.previous is None and w.following.name == "EE_SAP_Release_2026Q2.Jun"

    def test_the_latest_release_is_open_until_today(self):
        w = lk.window_for(self.rels, "EE_SAP_Release_2026Q3.Sept", TODAY)
        assert (w.end, w.open_ended) == (TODAY, True)
        assert w.previous.name == "EE_SAP_Release_2026Q2.Jun"
        assert w.following.name == "EE_SAP_Release_2026Q4.Oct"      # planned, not shipped

    def test_an_unreleased_version_has_no_window(self):
        assert lk.window_for(self.rels, "EE_SAP_Release_2026Q4.Oct", TODAY) is None


# ── which incidents count ────────────────────────────────────────────────────
class TestDeliverysRules:
    def test_the_exclusions_are_deliverys_jql(self):
        assert lk.EXCLUDED_ROOT_CAUSES == {
            "requirement/documentation", "new requirement", "non reproducible",
            "expected behaviour", "data issue", "duplicate", "security issue",
            "not applicable", "not a bug"}
        assert lk.EXCLUDED_COMPONENTS == {"ios app", "android app"}
        assert lk.EXCLUDED_LABELS == {"mapp", "ios", "android"}

    @pytest.mark.parametrize(("change", "reason"), [
        ({"status": "Cancelled"}, "Cancelled"),
        ({"rc": ""}, "No root cause yet"),
        ({"rc": "Data issue"}, "Root cause: Data issue"),
        ({"rc": "NOT A BUG"}, "Root cause: NOT A BUG"),
        ({"components": ["Ios App"]}, "App component: Ios App"),
        ({"labels": ["MAPP"]}, "App label: MAPP"),
    ])
    def test_each_exclusion_says_why(self, change, reason):
        row = _incident("X-1", None, **change)
        assert lk.exclusion(row) == reason

    def test_web_incidents_with_an_analysed_root_cause_count(self):
        for rc in ("Code issue", "Configuration", "3rd party issue", "connectivity issue"):
            assert lk.exclusion(_incident("X-1", None, rc=rc, components=["Checkout"])) is None

    def test_split_keeps_every_excluded_incident_visible(self):
        rows = [_incident("A-1", None), _incident("A-2", None, rc="duplicate")]
        counted, excluded = lk.split(rows)
        assert [r["key"] for r in counted] == ["A-1"]
        assert excluded[0]["excluded_because"] == "Root cause: duplicate"


class TestUatIssues:
    """Delivery's denominator, verified on 2026-09-28: every Bug of the
    fixVersion (cancelled ones too), and the Defects created since the UAT
    start — for which the previous release's date stands in."""

    SINCE = date(2026, 3, 16)                    # EE_SAP_Release_2026Q1.Mar shipped

    @staticmethod
    def _issue(kind, day, status="Done"):
        return {"key": f"{kind}-{day}", "type": kind, "status": status,
                "created": datetime.combine(day, datetime.min.time(), timezone.utc)}

    def test_every_bug_counts_whatever_its_date_or_status(self):
        rows = [self._issue("Bug", date(2025, 12, 1)),
                self._issue("Bug", date(2026, 4, 1), status="Cancelled")]
        counted, early = lk.split_uat(rows, self.SINCE)
        assert len(counted) == 2 and early == []

    def test_a_defect_older_than_the_previous_release_is_left_out_and_said_so(self):
        rows = [self._issue("Defect", date(2026, 3, 1)), self._issue("Defect", date(2026, 3, 16))]
        counted, early = lk.split_uat(rows, self.SINCE)
        assert [r["created"].date() for r in counted] == [date(2026, 3, 16)]
        assert "before the previous release (2026-03-16)" in early[0]["excluded_because"]

    def test_without_a_previous_release_every_defect_counts(self):
        counted, early = lk.split_uat([self._issue("Defect", date(2020, 1, 1))], None)
        assert len(counted) == 1 and early == []


class TestRatios:
    def test_ee_april_as_delivery_counted_it(self):
        """EE_SAP_Release_2026Q2.Apr in Delivery's Q2 workbook: 77 Bugs (17 of
        them cancelled) + 12 Defects = 89 UAT issues; 12 leaked, 6 High/Highest:
        13.5% and 6.7%.  Jira holds 13 Defects: one predates the previous
        release and is left out, exactly as the report left it out."""
        since = date(2026, 3, 16)
        at = lambda d: datetime.combine(d, datetime.min.time(), timezone.utc)
        rows = ([{"type": "Bug", "status": "Done", "created": at(date(2026, 3, 20))}] * 60
                + [{"type": "Bug", "status": "Cancelled", "created": at(date(2026, 1, 5))}] * 17
                + [{"type": "Defect", "status": "Done", "created": at(date(2026, 4, 2))}] * 12
                + [{"type": "Defect", "status": "Done", "created": at(date(2026, 3, 2))}])
        uat, early = lk.split_uat(rows, since)
        rel = lk.Release("EE_SAP_Release_2026Q2.Apr", True, date(2026, 4, 27))
        leaks = ([_incident(f"L-{i}", None, priority="High") for i in range(6)]
                 + [_incident(f"M-{i}", None) for i in range(6)])
        d = lk.ReleaseLeakage(EE, rel, None, uat, leaks, uat_left_out=early)
        assert (len(d.uat), d.uat_bugs, d.uat_defects, len(early)) == (89, 77, 12, 1)
        assert (len(d.leaks), d.leaks_hh) == (12, 6)
        assert f"{d.ratio:.1%}" == "13.5%"
        assert f"{d.ratio_hh:.1%}" == "6.7%"          # over ALL UAT issues

    def test_no_uat_issue_means_no_ratio(self):
        assert lk.ratio(3, 0) is None


class TestJiraRows:
    def test_an_incident_is_normalised(self):
        ids = {"Root Cause (EU)": "cf_rc", "ENVIRONMENT": "cf_env",
               "Steps to Reproduce": "cf_steps", "TestRail Case ID": "cf_tr"}
        issue = {"key": "EE20-9", "fields": {
            "summary": "Voucher rejected", "created": "2026-05-02T08:15:00.000+0200",
            "issuetype": {"name": "Production Incident"}, "priority": {"name": "High"},
            "status": {"name": "Closed"}, "components": [{"name": "Checkout"}],
            "labels": ["website"], "cf_rc": {"value": "Code issue"},
            "cf_env": [{"value": "PROD"}], "cf_tr": "C123456",
            "description": {"type": "doc", "content": [{"type": "paragraph", "content": [
                {"type": "text", "text": "Voucher VIP10 fails"}]}]},
            "cf_steps": {"type": "doc", "content": [{"type": "orderedList", "content": [
                {"type": "listItem", "content": [{"type": "paragraph", "content": [
                    {"type": "text", "text": "Add item"}]}]}]}]},
            "issuelinks": [{"type": {"outward": "clones"}, "outwardIssue": {
                "key": "EE20-1", "fields": {"summary": "Story", "issuetype": {"name": "Story"}}}}],
        }}
        r = lk.normalise_incident(issue, "https://j", ids)
        assert (r["root_cause"], r["environment"], r["priority"], r["case_refs"]) == (
            "Code issue", "PROD", "High", "C123456")
        assert r["description"] == "Voucher VIP10 fails"
        assert r["steps"] == "1. Add item"
        assert r["links"] == [{"key": "EE20-1", "relation": "clones", "type": "Story",
                               "summary": "Story"}]
        assert r["created"] == datetime(2026, 5, 2, 8, 15, tzinfo=timezone.utc)


class TestJiraReadsEverything:
    """A truncated count is a wrong count: the search pages until Jira says
    it is done, and a failure raises instead of reading as zero."""

    class _Resp:
        def __init__(self, payload, ok=True, status=200):
            self._payload, self.ok, self.status_code = payload, ok, status

        def json(self):
            return self._payload

    def test_every_page_is_read(self, monkeypatch):
        pages = [
            {"issues": [{"key": f"A-{i}"} for i in range(100)], "nextPageToken": "t1"},
            {"issues": [{"key": f"B-{i}"} for i in range(40)], "isLast": True},
        ]
        bodies = []

        def _post(url, json, auth, timeout):
            bodies.append(json)
            return self._Resp(pages[len(bodies) - 1])

        monkeypatch.setattr(jc, "_conf", lambda: ("https://j", "u", "t"))
        monkeypatch.setattr(jc.requests, "post", _post)
        assert len(jc.search_all("issuetype = x", ("project",))) == 140
        assert "nextPageToken" not in bodies[0] and bodies[1]["nextPageToken"] == "t1"

    def test_a_failed_page_raises_rather_than_undercounting(self, monkeypatch):
        monkeypatch.setattr(jc, "_conf", lambda: ("https://j", "u", "t"))
        monkeypatch.setattr(jc.requests, "post",
                            lambda *a, **k: self._Resp({}, ok=False, status=503))
        with pytest.raises(RuntimeError, match="503"):
            jc.search_all("issuetype = x", ("project",))

    def test_versions_that_cannot_be_read_raise(self, monkeypatch):
        monkeypatch.setattr(jc, "_conf", lambda: ("https://j", "u", "t"))
        monkeypatch.setattr(jc.requests, "get",
                            lambda *a, **k: self._Resp({}, ok=False, status=500))
        jc.project_versions.clear()
        with pytest.raises(RuntimeError, match="500"):
            jc.project_versions("EE20")

    def test_keys_are_found_bare_in_urls_and_lowercase(self):
        text = ("IPXL20-15740, https://x.atlassian.net/browse/SD-512 and "
                "ipxl20-15740 again\nhttps://x.atlassian.net/jira/software/c/"
                "projects/TPS/boards/1?selectedIssue=TPS-9")
        assert jc.extract_issue_keys(text) == ["IPXL20-15740", "SD-512", "TPS-9"]


# ── TestRail matching ────────────────────────────────────────────────────────
def _index():
    cases = pd.DataFrame([
        {"case_id": 101, "title": "Apply voucher code at checkout", "section_path": "Checkout > Vouchers",
         "url": "u101", "refs": "EE20-1", "labels": ["big_regr_desktop"]},
        {"case_id": 102, "title": "Pay with credit card", "section_path": "Checkout > Payment",
         "url": "u102", "refs": "", "labels": ["big_regr_desktop"]},
        {"case_id": 103, "title": "Search returns products", "section_path": "Search",
         "url": "u103", "refs": "EE20-9", "labels": []},
        {"case_id": 104, "title": "Remove voucher from basket", "section_path": "Basket > Vouchers",
         "url": "u104", "refs": "", "labels": ["big_regr_mobile"]},
    ])
    return lm.build_index(cases, automated={101, 104}, outdated={104})


class TestMatching:
    def test_statuses_follow_the_backlog(self):
        idx = _index()
        assert (idx.cases[101].status, idx.cases[102].status,
                idx.cases[103].status, idx.cases[104].status) == (
            lm.AUTOMATED, lm.MANUAL, lm.MANUAL_OUT, lm.OUTDATED)

    def test_direct_links_come_first_and_say_why(self):
        idx = _index()
        row = _incident("EE20-9", None)
        row.update(case_refs="TestRail Case ID: C102",
                   links=[{"key": "EE20-1", "relation": "clones", "type": "Story", "summary": ""}])
        links = {c.case_id: why for c, why in lm.direct_links(row, idx)}
        assert set(links) == {103, 102, 101}
        assert "cites EE20-9" in links[103]
        assert "TestRail Case ID" in links[102]
        assert "clones" in links[101]

    def test_candidates_rank_by_shared_distinctive_words(self):
        idx = _index()
        row = _incident("EE20-50", None)
        row.update(summary="Voucher not applied at checkout")
        cands = [c.case_id for c in lm.candidates(row, idx, exclude=set())]
        assert cands[0] == 101 and 104 in cands and 103 not in cands

    def test_linked_cases_are_not_offered_twice(self):
        idx = _index()
        row = _incident("EE20-50", None)
        row.update(summary="Voucher not applied at checkout")
        assert 101 not in [c.case_id for c in lm.candidates(row, idx, exclude={101})]

    @pytest.mark.parametrize(("statuses", "expected"), [
        ([], "No test case"),
        ([lm.AUTOMATED], "Automated test"),
        ([lm.OUTDATED], "Automated test outdated"),
        ([lm.MANUAL], "Manual test"),
        ([lm.MANUAL_OUT], "Manual, not in regression"),
        ([lm.MANUAL, lm.AUTOMATED], "Manual and automated"),
    ])
    def test_the_gap_follows_the_covering_case(self, statuses, expected):
        assert lm.gap(statuses) == expected


# ── the AI's answer ──────────────────────────────────────────────────────────
def _verdict(**over):
    v = {"key": "EE20-50", "uat_detectability": lai.DETECTABLE,
         "category": "Automated test miss", "functional_area": "Checkout > Vouchers",
         "confidence": 0.8, "rationale": "Voucher flow is testable in UAT.",
         "evidence": [{"field": "Summary", "value": "Voucher", "reason": "core flow"}],
         "missing_information": [], "corrective_actions": ["Review C101 assertions"],
         "testrail_matches": [{"case_id": 101, "confidence": 0.85, "reason": "voucher at checkout"}]}
    v.update(over)
    return v


class TestTheAIsAnswer:
    def test_a_good_verdict_proposes_the_likely_case(self):
        v = lai.validate(_verdict(), {101, 104})
        assert v["testrail_case"] == "C101" and not v["requires_review"] and not v["flags"]

    def test_a_case_it_was_not_shown_is_dropped(self):
        v = lai.validate(_verdict(testrail_matches=[{"case_id": 999, "confidence": 0.9,
                                                     "reason": "x"}]), {101})
        assert v["testrail_matches"] == [] and v["testrail_case"] == ""

    def test_a_low_confidence_match_is_never_the_covering_case(self):
        v = lai.validate(_verdict(testrail_matches=[{"case_id": 101, "confidence": 0.4,
                                                     "reason": "maybe"}]), {101})
        assert v["testrail_case"] == "" and v["testrail_matches"][0]["case_id"] == 101

    def test_an_unknown_category_is_flagged_not_trusted(self):
        v = lai.validate(_verdict(category="Bad luck"), {101})
        assert v["category"] == "Needs more information" and v["requires_review"]

    def test_a_category_of_the_other_verdict_is_flagged(self):
        v = lai.validate(_verdict(category="Production data"), {101})
        assert v["flags"] and v["requires_review"]

    def test_low_confidence_needs_a_person(self):
        assert lai.validate(_verdict(confidence=0.5), {101})["requires_review"]

    def test_every_category_belongs_to_a_verdict(self):
        assert set(lai.VERDICT_OF.values()) <= set(lai.VERDICTS)
        assert len(set(lai.CATEGORIES)) == len(lai.CATEGORIES)

    def test_the_prompt_forbids_the_production_shortcut(self):
        text = lai.system_instruction()
        assert "merely because the incident happened in" in text
        assert "production-only promotion" in text

    def test_the_block_shows_linked_cases_and_candidates(self):
        idx = _index()
        row = _incident("EE20-50", None)
        block = lai.incident_block(row, [(idx.cases[101], "cites it")], [idx.cases[102]])
        assert "LINKED (cites it)" in block and "C102 [Manual, in regression]" in block

    def test_the_schema_is_accepted_by_the_sdk(self):
        types = pytest.importorskip("google.genai.types")
        types.GenerateContentConfig(response_mime_type="application/json",
                                    response_schema=lai.response_schema())

    def test_classify_reads_the_json_and_keeps_only_asked_keys(self, monkeypatch):
        from src import gemini_client as gc
        idx = _index()
        rows = [_incident("EE20-50", None)]
        answer = '```json\n{"verdicts": [' + __import__("json").dumps(_verdict()) + \
                 ', {"key": "EE20-999", "category": "Production data"}]}\n```'
        monkeypatch.setattr(gc, "ready", lambda: True)
        monkeypatch.setattr(gc, "types", pytest.importorskip("google.genai.types"))
        monkeypatch.setattr(gc, "generate", lambda *a, **k: gc.Result(answer, "gemini-3.8-flash"))
        out, model, errors = lai.classify(rows, {"EE20-50": ([], [idx.cases[101]])})
        assert set(out) == {"EE20-50"} and model == "gemini-3.8-flash" and not errors

    def test_when_every_model_refuses_it_says_so(self, monkeypatch):
        from src import gemini_client as gc
        monkeypatch.setattr(gc, "ready", lambda: True)
        monkeypatch.setattr(gc, "types", pytest.importorskip("google.genai.types"))
        monkeypatch.setattr(gc, "generate", lambda *a, **k: gc.Result(None, None, "429"))
        out, model, errors = lai.classify([_incident("A-1", None)], {"A-1": ([], [])})
        assert out == {} and model is None and errors


# ── the Key QA's review ──────────────────────────────────────────────────────
class TestStore:
    def test_a_review_never_erases_the_ai_proposal(self):
        s = ls.MemoryStore()
        s.put_ai("EE20", "A-1", {"uat_detectability": lai.DETECTABLE,
                                 "category": "Missing test case", "testrail_case": ""}, "m", "h")
        s.review("EE20", "A-1", "Anna", ls.CHANGED,
                 {"uat_detectability": lai.NOT_DETECTABLE, "category": "Production data",
                  "testrail_case": ""}, "promo only in prod")
        rec = s.get("EE20", "A-1")
        assert rec.ai["uat_detectability"] == lai.DETECTABLE
        assert rec.final["uat_detectability"] == lai.NOT_DETECTABLE
        assert (rec.status, rec.reviewer) == (ls.CHANGED, "Anna")
        assert rec.events[0].before["category"] == "Missing test case"

    def test_a_re_analysis_keeps_the_previous_proposal(self):
        s = ls.MemoryStore()
        s.put_ai("EE20", "A-1", {"category": "one"}, "m1", "h1")
        s.put_ai("EE20", "A-1", {"category": "two"}, "m2", "h2")
        rec = s.get("EE20", "A-1")
        assert rec.ai["category"] == "two" and rec.past_ai[0]["ai"]["category"] == "one"

    def test_the_record_is_a_copy(self):
        s = ls.MemoryStore()
        s.put_ai("EE20", "A-1", {"category": "one"}, "m", "h")
        s.get("EE20", "A-1").ai["category"] = "tampered"
        assert s.get("EE20", "A-1").ai["category"] == "one"

    def test_attempts_are_remembered_until_forgotten(self):
        s = ls.MemoryStore()
        s.mark_attempted("EE20", "R1")
        assert s.attempted("EE20", "R1")
        s.forget_attempt("EE20", "R1")
        assert not s.attempted("EE20", "R1")


class TestExport:
    def test_the_workbook_has_every_sheet_and_keeps_ai_and_human_apart(self):
        import openpyxl
        rel = lk.Release("EE_SAP_Release_2026Q2.Apr", True, date(2026, 4, 27))
        created = datetime(2026, 5, 2, tzinfo=timezone.utc)
        leak = _incident("A-1", created)
        uat = [{"key": "B-1", "url": "u", "type": "Bug", "created": created,
                "priority": "High", "status": "Done", "summary": "s"}]
        excluded = [{**_incident("A-2", created, rc="duplicate"), "excluded_because": "Root cause: duplicate"}]
        d = lk.ReleaseLeakage(EE, rel, None, uat, [leak], excluded)
        s = ls.MemoryStore()
        s.put_ai("EE20", "A-1", lai.validate(_verdict(key="A-1"), {101}), "gemini-3.8-flash", "h")
        s.review("EE20", "A-1", "Anna", ls.CONFIRMED, s.get("EE20", "A-1").final)
        data = lx.workbook(d, {"A-1": s.get("EE20", "A-1")}, {"A-1": "Automated test"}, {})
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Summary", "Incidents", "UAT issues", "Excluded", "Audit trail"]
        header = [c.value for c in wb["Incidents"][1]]
        assert {"AI verdict", "AI rationale", "Final verdict", "Reviewer",
                "Coverage gap"} <= set(header)
        assert wb["Audit trail"].max_row == 2


# ── the tab ──────────────────────────────────────────────────────────────────
class TestTab:
    @staticmethod
    def _run(bu):
        from streamlit.testing.v1 import AppTest

        def page(bu):
            from datetime import date, datetime, timezone

            import pandas as pd

            from src import gemini_client, jira_client
            from src import leakage as lk
            from src import leakage_match as lm
            from src.ui import global_filter, leakage_tab

            created = datetime(2026, 9, 23, tzinfo=timezone.utc)
            rels = [lk.Release("EE_SAP_Release_2026Q3.Jul", True, date(2026, 7, 22)),
                    lk.Release("EE_SAP_Release_2026Q3.Sept", True, date(2026, 9, 21)),
                    lk.Release("EE_SAP_Release_2026Q4.Oct", False, date(2026, 10, 7))]

            def row(key, **kw):
                r = {"key": key, "url": f"https://j/browse/{key}", "type": "Production Incident",
                     "summary": "Voucher fails", "created": created, "priority": "High",
                     "status": "Done", "components": ["Checkout"], "labels": [],
                     "root_cause": "Code issue", "environment": "PROD", "description": "",
                     "steps": "", "actual": "", "expected": "", "case_refs": "", "links": [],
                     "resolution": ""}
                r.update(kw)
                return r

            def analyse(group, rels_, name, today):
                win = lk.window_for(rels_, name, date(2026, 9, 28))
                uat = [{"key": f"B-{i}", "url": "u", "type": "Bug", "created": created,
                        "priority": "Medium", "status": "Done", "summary": "s"} for i in range(8)]
                return lk.ReleaseLeakage(group, rels_[1], win, uat, [row("EE20-1")],
                                         [{**row("EE20-2", root_cause=""),
                                           "excluded_because": "No root cause yet"}])

            index = lm.build_index(pd.DataFrame([{
                "case_id": 101, "title": "Apply voucher", "section_path": "Checkout",
                "url": "https://t/index.php?/cases/view/101", "refs": "", "labels": []}]),
                {101}, set())
            saved = (jira_client.available, lk.releases, lk.analyse, lm.index_for,
                     global_filter.current, gemini_client.ready)
            jira_client.available = lambda: True
            lk.releases = lambda g: rels
            lk.analyse = analyse
            lm.index_for = lambda bus: index
            global_filter.current = lambda: ("website", bu)
            gemini_client.ready = lambda: False
            try:
                leakage_tab.render()
            finally:
                (jira_client.available, lk.releases, lk.analyse, lm.index_for,
                 global_filter.current, gemini_client.ready) = saved

        at = AppTest.from_function(page, args=(bu,), default_timeout=30)
        at.run()
        assert not at.exception, at.exception
        return at

    def test_the_latest_release_is_analysed_by_default(self):
        at = self._run("Drogas")
        text = " ".join(m.value for m in at.markdown) + " ".join(c.value for c in at.caption)
        assert "12.5%" in text                           # 1 leaked ÷ 8 UAT issues
        assert "EE_SAP_Release_2026Q3.Sept" in text
        assert "until EE_SAP_Release_2026Q4.Oct ships" in text
        assert at.selectbox[0].value == "EE_SAP_Release_2026Q3.Sept"

    def test_a_shared_project_says_it_cannot_be_split(self):
        at = self._run("Watsons Turkey")
        assert any("cannot be split per Business Unit" in c.value for c in at.caption)

    def test_release_names_lose_only_the_part_they_share(self):
        from src.ui.leakage_tab import short_names
        assert short_names(["EE_SAP_Release_2026Q2.Apr", "EE_SAP_Release_2026Q3.Sept"]) == {
            "EE_SAP_Release_2026Q2.Apr": "2026Q2.Apr", "EE_SAP_Release_2026Q3.Sept": "2026Q3.Sept"}
        assert set(short_names(["KV_SAP_Release_2026Q2.1", "KV_SAP_Release_2026Q2.2"]).values()) == {
            "2026Q2.1", "2026Q2.2"}                         # never cut inside a token
        assert short_names(["SD_Release_2026Q2.Apr SAP", "SD_Release_2026Q2.May SAP"])[
            "SD_Release_2026Q2.May SAP"] == "2026Q2.May"

    def test_without_ai_the_counts_still_stand(self):
        at = self._run("Drogas")
        assert any("GEMINI_API_KEY" in i.value for i in at.info)
