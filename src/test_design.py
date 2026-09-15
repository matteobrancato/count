"""AI Test Design — user stories, pages and documents in, TestRail-style test
cases out.

Pure logic, no widgets: reading the sources, building the one Gemini request,
and parsing + validating what comes back.  The tab (`ui/test_design_tab.py`)
only arranges it on screen.

The design goal is a SMALL, PRECISE set of test cases traceable to the
acceptance criteria — not a long list.  Three things hold it there:
  * the model must first extract the AC, and every test must name the AC it
    covers, so an extra test with nothing to cover is visible at a glance;
  * the model answers in a fixed JSON schema, never free text, so what is shown
    is exactly what was asked for (and phase 2 can write it to TestRail as is);
  * ambiguities go into "open questions" instead of being resolved by guessing.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from xml.etree import ElementTree

from . import gemini_client

# Best model first, then step down when Google refuses one (see gemini_client).
# Pro leads because writing tests from AC is reasoning, not lookup; on the free
# tier it often has little or no quota, and the chain then moves on by itself.
MODEL_CHAIN: list[str] = [
    "gemini-2.5-pro",
    "gemini-2.5-flash",
    "gemini-2.5-flash-lite",
]

MAX_FILES = 10
# Gemini's inline request limit is ~20 MB; the rest is headroom for the text.
MAX_TOTAL_BYTES = 15 * 1024 * 1024
# Per text source — enough for a long spec, bounded so one document cannot
# crowd out the stories it is meant to support.
MAX_TEXT_CHARS = 60_000

IMAGE_MIME = {"png": "image/png", "jpg": "image/jpeg",
              "jpeg": "image/jpeg", "webp": "image/webp"}
TEXT_EXT = {"txt", "md", "csv", "json", "feature", "html", "xml"}
UPLOAD_TYPES = sorted({*IMAGE_MIME, "pdf", "docx", "xlsx", *TEXT_EXT})

PRIORITIES = ("Highest", "High", "Medium", "Low")
DEVICES    = ("Both", "Desktop", "Mobile")


# ── sources ──────────────────────────────────────────────────────────────────
@dataclass
class Source:
    """One thing the AI reads, and what the user is told about it."""
    kind: str                     # "jira" | "confluence" | "file" | "notes"
    label: str                    # shown to the user
    ok: bool                      # False → reported, never sent
    note: str = ""                # "acceptance criteria found", "not found", …
    text: str = ""                # text sent to the model
    blob: bytes | None = None     # sent natively (PDF, images)
    mime: str | None = None
    warn: bool = False            # read, but something the user should know


def _clip(text: str) -> tuple[str, bool]:
    text = (text or "").strip()
    if len(text) <= MAX_TEXT_CHARS:
        return text, False
    return text[:MAX_TEXT_CHARS], True


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_text(data: bytes) -> str:
    """Paragraphs and tables of a .docx, in document order — no dependency."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ElementTree.fromstring(z.read("word/document.xml"))
    body = root.find(f"{_W}body")
    if body is None:
        return ""

    def para(p) -> str:
        return "".join(t.text or "" for t in p.iter(f"{_W}t"))

    lines: list[str] = []
    for el in body:
        if el.tag == f"{_W}p":
            lines.append(para(el))
        elif el.tag == f"{_W}tbl":
            for row in el.iter(f"{_W}tr"):
                cells = [" ".join(para(p) for p in c.iter(f"{_W}p")).strip()
                         for c in row.iter(f"{_W}tc")]
                lines.append(" | ".join(cells))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def xlsx_text(data: bytes, max_rows: int = 2000) -> str:
    """Every sheet as `a | b | c` rows (openpyxl is already a dependency)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                out.append(f"(… truncated after {max_rows} rows)")
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                out.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(out).strip()


def read_file(name: str, data: bytes) -> Source:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    try:
        if ext in IMAGE_MIME:
            return Source("file", name, True, "image", blob=data, mime=IMAGE_MIME[ext])
        if ext == "pdf":
            return Source("file", name, True, "PDF", blob=data, mime="application/pdf")
        if ext == "docx":
            raw = docx_text(data)
        elif ext == "xlsx":
            raw = xlsx_text(data)
        elif ext in TEXT_EXT:
            raw = data.decode("utf-8-sig", errors="replace")
        else:
            return Source("file", name, False, f"unsupported file type .{ext}")
    except Exception as exc:                                            # noqa: BLE001
        return Source("file", name, False, f"could not be read ({type(exc).__name__})")
    text, clipped = _clip(raw)
    if not text:
        return Source("file", name, False, "no readable text")
    note = f"{len(text):,} characters" + (" · truncated" if clipped else "")
    return Source("file", name, True, note, text=text, warn=clipped)


_AC_IN_TEXT = re.compile(
    r"acceptance\s+criteria|\bAC\s*\d|\bgiven\b[\s\S]{0,200}?\bwhen\b[\s\S]{0,200}?\bthen\b",
    re.IGNORECASE)


def story_source(story: dict) -> Source:
    key = story.get("key", "")
    if story.get("error"):
        return Source("jira", key, False, story["error"])
    label = f"{key} — {story.get('summary') or '(no summary)'}"
    ac, desc = story.get("acceptance_criteria", ""), story.get("description", "")
    if ac:
        note, warn = "acceptance criteria field", False
    elif _AC_IN_TEXT.search(desc):
        note, warn = "acceptance criteria in the description", False
    else:
        note, warn = "no acceptance criteria found — tests will be derived", True
    if story.get("attachments"):
        note += f" · {story['attachments']} attachment(s) not read"
    body = (f"Jira {story.get('type') or 'issue'} {key}: {story.get('summary', '')}\n"
            f"Status: {story.get('status', '')}\n\n"
            f"ACCEPTANCE CRITERIA FIELD:\n{ac or '(empty)'}\n\n"
            f"DESCRIPTION:\n{desc or '(empty)'}")
    text, clipped = _clip(body)
    return Source("jira", label, True, note + (" · truncated" if clipped else ""),
                  text=text, warn=warn or clipped)


def page_source(page: dict) -> Source:
    ref = page.get("ref", "")
    if page.get("error"):
        return Source("confluence", f"Page {ref}", False, page["error"])
    label = f"{page.get('title') or 'Untitled'} ({page.get('id')})"
    text, clipped = _clip(f"Confluence page: {page.get('title', '')}\n\n{page.get('text', '')}")
    note = f"{len(page.get('text', '')):,} characters" + (" · truncated" if clipped else "")
    return Source("confluence", label, True, note, text=text, warn=clipped)


def notes_source(notes: str) -> Source | None:
    text, clipped = _clip(notes)
    if not text:
        return None
    return Source("notes", "Your notes", True,
                  f"{len(text):,} characters" + (" · truncated" if clipped else ""),
                  text=text, warn=clipped)


# ── the request ──────────────────────────────────────────────────────────────
SYSTEM_INSTRUCTION = """
You are a senior QA engineer at AS Watson writing MANUAL test cases for
TestRail, template "Test Case (Steps)". You work ONLY from the material you are
given: Jira stories, Confluence pages, documents, UI images and the user's notes.

1. ACCEPTANCE CRITERIA
   - Extract the acceptance criteria (AC) of every story: the acceptance-criteria
     field, or an "Acceptance criteria"/"AC"/Given-When-Then section.
   - Only when a source has NO explicit AC, derive them from clearly stated
     required behaviour and start their text with "(derived)".
   - Number them AC1..ACn in source order. Keep the wording faithful and short.
     "source" is the Jira key, page title or file name it came from.

2. THE SMALLEST SET OF TEST CASES THAT COVERS EVERY AC
   - Default: one test case per AC. Merge ACs that are verified in one flow
     with the same preconditions. Split only when preconditions or expected
     outcomes genuinely differ.
   - Never test behaviour that no AC states. No generic tests (login works,
     page loads, cross-browser, performance) unless an AC requires them.
   - Negative and edge cases: {negative_rule}
   - Device scope: {device_rule}

3. EACH TEST CASE
   - title: one sentence naming the behaviour under test, starting with the
     actor, e.g. "User applies a valid promo code at checkout".
   - preconditions: the concrete state and data needed (user type, basket
     content, configuration, feature flag). Empty string if none.
   - steps: each step is ONE user action in the imperative ("Open the basket
     page"), with the observable result expected right after it. Usually 3 to
     10 steps. No step without an expected result. No vague results ("works
     correctly") — state what the user sees.
   - priority: Highest = purchase, payment, legal or security flow; High = the
     main behaviour of the story; Medium = secondary behaviour; Low = cosmetic.
   - type: "Functional", unless a source says otherwise.
   - covers: the AC ids this test verifies. Every AC must be covered.

4. OPEN QUESTIONS
   - List ambiguities, contradictions, missing data and unmeasurable wording in
     the AC that prevent a precise test (e.g. "'quickly' is not measurable").
   - NEVER resolve them by guessing. Empty list if there are none.

5. RULES
   - Write in English.
   - Never invent URLs, credentials, prices, product names or test data that
     are not in the sources: use placeholders such as <valid promo code>.
   - Images are UI designs or screenshots: use them to name visible elements
     precisely, not as extra requirements.
   - The user's notes are both context and instructions; follow them unless
     they contradict the rules above.
""".strip()

_NEGATIVE = {
    False: ("only when an AC explicitly states a validation, error, limit or "
            "rejection — then test that behaviour. Otherwise none."),
    True:  ("add the few high-value negative and edge cases the AC imply "
            "(invalid input, limits, error states), each linked to its AC. "
            "Keep them few."),
}
_DEVICE = {
    "Both":    ("Desktop and Mobile. Write ONE test per behaviour with device "
                "\"Both\"; split by device only when an AC differs by device."),
    "Desktop": "Desktop only. Every test has device \"Desktop\".",
    "Mobile":  "Mobile only. Every test has device \"Mobile\".",
}


def system_instruction(device: str, include_negative: bool) -> str:
    return (SYSTEM_INSTRUCTION
            .replace("{negative_rule}", _NEGATIVE[bool(include_negative)])
            .replace("{device_rule}", _DEVICE.get(device, _DEVICE["Both"])))


def request_parts(sources: list[Source]) -> list[tuple]:
    """The request body as plain tuples — ("text", str) or ("bytes", data, mime).

    Kept free of SDK types so it can be tested without Gemini installed.
    """
    usable = [s for s in sources if s.ok]
    intro = (f"Design test cases from the {len(usable)} source(s) below. "
             "Answer with the JSON object described by the schema.")
    parts: list[tuple] = [("text", intro)]
    for s in usable:
        header = f"===== SOURCE ({s.kind}): {s.label} ====="
        if s.blob is not None:
            parts.append(("text", header))
            parts.append(("bytes", s.blob, s.mime))
        else:
            parts.append(("text", f"{header}\n{s.text}"))
    return parts


def response_schema():
    """The answer's shape, as a pydantic model the SDK turns into a schema."""
    from typing import Literal

    from pydantic import BaseModel

    class Step(BaseModel):
        step: str
        expected: str

    class Criterion(BaseModel):
        id: str
        text: str
        source: str

    class TestCase(BaseModel):
        title: str
        preconditions: str
        priority: Literal["Highest", "High", "Medium", "Low"]
        type: str
        device: Literal["Both", "Desktop", "Mobile"]
        covers: list[str]
        steps: list[Step]

    class Design(BaseModel):
        acceptance_criteria: list[Criterion]
        test_cases: list[TestCase]
        open_questions: list[str]

    return Design


# ── the answer ───────────────────────────────────────────────────────────────
class DesignError(ValueError):
    """The model answered, but not with a usable design."""


@dataclass
class TestCase:
    id: str
    title: str
    preconditions: str
    priority: str
    type: str
    device: str
    covers: list[str]
    steps: list[tuple[str, str]]


@dataclass
class Design:
    criteria: list[dict]                    # {"id", "text", "source"}
    tests: list[TestCase]
    questions: list[str]
    uncovered: list[str] = field(default_factory=list)
    dropped: int = 0                        # tests discarded for having no steps

    def tests_for(self, ac_id: str) -> list[str]:
        return [t.id for t in self.tests if ac_id in t.covers]


def _load_json(text: str) -> dict:
    raw = (text or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw)
    try:
        return json.loads(raw)
    except ValueError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            try:
                return json.loads(raw[start:end + 1])
            except ValueError:
                pass
    raise DesignError("the model's answer was not valid JSON")


def _s(v) -> str:
    return str(v).strip() if v is not None else ""


def parse_design(text: str) -> Design:
    """Validate the model's JSON into a Design the screen can trust.

    Ids are renumbered (AC1…, TC1…) so what is shown is always consistent,
    references to unknown AC are dropped, a test with no usable step is
    discarded and counted, and every AC no test covers is listed.
    """
    data = _load_json(text)
    if not isinstance(data, dict):
        raise DesignError("the model's answer had the wrong shape")

    criteria: list[dict] = []
    renumber: dict[str, str] = {}
    for i, c in enumerate(data.get("acceptance_criteria") or [], 1):
        if not isinstance(c, dict) or not _s(c.get("text")):
            continue
        new_id = f"AC{len(criteria) + 1}"
        renumber[_s(c.get("id")) or f"#{i}"] = new_id
        criteria.append({"id": new_id, "text": _s(c.get("text")),
                         "source": _s(c.get("source"))})

    tests: list[TestCase] = []
    dropped = 0
    for t in data.get("test_cases") or []:
        if not isinstance(t, dict):
            continue
        steps = [(_s(s.get("step")), _s(s.get("expected")))
                 for s in (t.get("steps") or [])
                 if isinstance(s, dict) and _s(s.get("step"))]
        if not steps or not _s(t.get("title")):
            dropped += 1
            continue
        covers = []
        for ref in t.get("covers") or []:
            new = renumber.get(_s(ref))
            if new and new not in covers:
                covers.append(new)
        priority = _s(t.get("priority")).title()
        device = _s(t.get("device")).title()
        tests.append(TestCase(
            id=f"TC{len(tests) + 1}",
            title=_s(t.get("title")),
            preconditions=_s(t.get("preconditions")),
            priority=priority if priority in PRIORITIES else "Medium",
            type=_s(t.get("type")) or "Functional",
            device=device if device in DEVICES else "Both",
            covers=covers,
            steps=steps,
        ))

    if not tests:
        raise DesignError("the model returned no usable test case")

    covered = {a for t in tests for a in t.covers}
    questions = [_s(q) for q in (data.get("open_questions") or []) if _s(q)]
    return Design(criteria=criteria, tests=tests, questions=questions,
                  uncovered=[c["id"] for c in criteria if c["id"] not in covered],
                  dropped=dropped)


def to_text(design: Design) -> str:
    """The whole design as plain text, for copying into TestRail or a ticket."""
    out: list[str] = []
    for t in design.tests:
        out.append(f"{t.id} · {t.title}")
        out.append(f"Priority: {t.priority} | Type: {t.type} | Device: {t.device}"
                   + (f" | Covers: {', '.join(t.covers)}" if t.covers else ""))
        out.append("Preconditions: " + (t.preconditions or "—"))
        out.append("Steps:")
        for n, (step, expected) in enumerate(t.steps, 1):
            out.append(f"  {n}. {step}")
            out.append(f"     Expected: {expected or '—'}")
        out.append("")
    if design.questions:
        out.append("Open questions:")
        out.extend(f"  - {q}" for q in design.questions)
    return "\n".join(out).strip()


# ── generation ───────────────────────────────────────────────────────────────
@dataclass
class Outcome:
    design: Design | None
    model: str | None
    error: str = ""
    raw: str = ""


def generate(sources: list[Source], device: str, include_negative: bool,
             cooling: dict[str, float]) -> Outcome:
    """One Gemini call: every usable source in, a validated Design out."""
    if not any(s.ok for s in sources):
        return Outcome(None, None, "None of the sources could be read.")
    types = gemini_client.types
    parts = []
    for p in request_parts(sources):
        if p[0] == "text":
            parts.append(types.Part.from_text(text=p[1]))
        else:
            parts.append(types.Part.from_bytes(data=p[1], mime_type=p[2]))
    config = types.GenerateContentConfig(
        system_instruction=system_instruction(device, include_negative),
        response_mime_type="application/json",
        response_schema=response_schema(),
        # Low, not zero: faithful to the AC, but free to phrase steps naturally.
        temperature=0.2,
    )
    result = gemini_client.generate(
        [types.Content(role="user", parts=parts)], config, MODEL_CHAIN, cooling)
    if result.model is None:
        return Outcome(None, None, gemini_client.failure_message(result.error))
    try:
        return Outcome(parse_design(result.text or ""), result.model, raw=result.text or "")
    except DesignError as exc:
        msg = str(exc)
        return Outcome(None, result.model, f"⚠️ {msg[:1].upper()}{msg[1:]}. Try again.",
                       raw=result.text or "")
